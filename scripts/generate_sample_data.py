"""生成合成样例数据（用于无教师数据时验证完整流程）。

生成内容：
- Word 制度文档（python-docx）
- PDF 制度文档（PyMuPDF）
- Excel 统计报表（openpyxl，含多级表头 / 合并单元格）
- 测试问题集（JSON）

输出到 data/documents、data/tables、data/questions。
未来拿到教师真实数据后，只需将文件放入对应目录即可运行，无需改动代码。
"""
from __future__ import annotations

import json
import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "data" / "documents"
TABLES_DIR = ROOT / "data" / "tables"
QUESTIONS_DIR = ROOT / "data" / "questions"
SAMPLE_DIR = ROOT / "data" / "sample"


# ---------------------------------------------------------------------------
# Word 文档内容
# ---------------------------------------------------------------------------
CAPITAL_RULES = [
    ("heading", "第一章 总则"),
    ("body", "第一条 为加强商业银行资本监管，维护银行体系稳健运行，制定本办法。"),
    ("body", "第二条 本办法适用于在中华人民共和国境内设立的商业银行。"),
    ("heading", "第二章 资本充足率要求"),
    ("body", "第三条 商业银行核心一级资本充足率不得低于5%。"),
    ("body", "第四条 商业银行一级资本充足率不得低于6%。"),
    ("body", "第五条 商业银行资本充足率不得低于8%。"),
    ("body", "第六条 商业银行杠杆率不得低于4%。"),
    ("body", "第七条 系统重要性银行应当额外满足附加资本要求，附加资本要求为1%。"),
    ("heading", "第三章 风险加权资产计量"),
    ("body", "第八条 商业银行应当采用权重法计量信用风险加权资产。"),
    ("body", "第九条 对一般企业债权的风险权重为100%。"),
    ("body", "第十条 对符合条件的微型和小型企业债权的风险权重为75%。"),
    ("heading", "第四章 监督检查"),
    ("body", "第十一条 商业银行应当每年开展一次内部资本充足评估。"),
    ("body", "第十二条 商业银行应当于每季度结束后30日内向监管机构报送资本充足率报表。"),
]

REPORT_RULES = [
    ("heading", "第一章 总则"),
    ("body", "第一条 为规范银行业金融机构重大事项报告工作，制定本制度。"),
    ("body", "第二条 本制度适用于政策性银行、商业银行、农村合作金融机构等。"),
    ("heading", "第二章 报告时限"),
    ("body", "第三条 银行业金融机构发生重大事项后，应当在5个工作日内向监管机构报告。"),
    ("body", "第四条 发生特别重大事项的，应当在24小时内先行电话报告，并在3个工作日内补报书面材料。"),
    ("body", "第五条 季度报告应当于每季度结束后20日内报送。"),
    ("body", "第六条 年度报告应当于年度结束后3个月内报送。"),
    ("heading", "第三章 报告内容"),
    ("body", "第七条 重大事项报告应当包括事项基本情况、影响分析、拟采取的处置措施。"),
    ("body", "第八条 涉及金额超过净资产10%的交易应当作为重大事项报告。"),
]

LOAN_RULES = [
    ("heading", "第一章 总则"),
    ("body", "第一条 为规范个人贷款业务经营行为，制定本办法。"),
    ("heading", "第二章 贷款条件"),
    ("body", "第二条 个人消费贷款期限最长不得超过10年。"),
    ("body", "第三条 个人经营贷款期限最长不得超过5年。"),
    ("body", "第四条 单笔贷款金额超过500万元的，应当追加有效担保。"),
    ("body", "第五条 贷款年利率不得违反国家规定的利率上限。"),
    ("heading", "第三章 风险管理"),
    ("body", "第六条 贷款人应当对借款人的还款能力进行尽职调查。"),
    ("body", "第七条 借款人连续3期未按约定偿还贷款的，贷款人应当启动贷后催收。"),
]


def _write_docx(path: Path, lines: list[tuple[str, str]]) -> None:
    from docx import Document

    doc = Document()
    for kind, text in lines:
        if kind == "heading":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    doc.save(str(path))


def _write_pdf(path: Path, lines: list[tuple[str, str]]) -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    y = 72.0
    for kind, text in lines:
        if kind == "heading":
            page.insert_text((72, y), text, fontsize=14, fontname="china-s")
            y += 26
        else:
            page.insert_text((72, y), text, fontsize=11, fontname="china-s")
            y += 20
        if y > 780:
            page = doc.new_page()
            y = 72.0
    doc.save(str(path))


# ---------------------------------------------------------------------------
# Excel 内容
# ---------------------------------------------------------------------------
def _write_indicator_table(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "年度数据"
    # 标题（合并）
    ws["A1"] = "商业银行主要监管指标统计表"
    ws.merge_cells("A1:G1")
    # 年份（多级表头，合并）
    ws["B2"] = "2023年"
    ws.merge_cells("B2:D2")
    ws["E2"] = "2024年"
    ws.merge_cells("E2:G2")
    # 季度表头
    headers = ["指标名称", "一季度", "二季度", "三季度", "一季度", "二季度", "三季度"]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=3, column=j, value=h)
    # 数据行
    data = [
        ("营业收入（亿元）", 120.5, 125.3, 130.8, 135.2, 140.6, 145.9),
        ("净利润（亿元）", 15.2, 16.8, 18.1, 19.5, 20.3, 21.7),
        ("不良贷款率（%）", 1.85, 1.82, 1.79, 1.75, 1.72, 1.68),
        ("资本充足率（%）", 13.5, 13.6, 13.7, 13.8, 13.9, 14.0),
    ]
    for i, row in enumerate(data, start=4):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    wb.save(str(path))


def _write_balance_table(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "机构数据"
    headers = ["机构", "指标", "2024年末", "2025年6月"]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=1, column=j, value=h)
    data = [
        ("A机构", "总资产（亿元）", 5000, 5200),
        ("A机构", "总负债（亿元）", 4600, 4780),
        ("A机构", "营业收入（亿元）", 300, 320),
        ("B机构", "总资产（亿元）", 8000, 8300),
        ("B机构", "总负债（亿元）", 7200, 7460),
        ("B机构", "营业收入（亿元）", 450, 480),
    ]
    for i, row in enumerate(data, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    wb.save(str(path))


# ---------------------------------------------------------------------------
# 测试问题集
# ---------------------------------------------------------------------------
QUESTIONS = [
    # ---- 文本题：条款 / 阈值 / 期限 ----
    {"question": "商业银行核心一级资本充足率不得低于多少？", "answer": "5%", "type": "text"},
    {"question": "商业银行一级资本充足率不得低于多少？", "answer": "6%", "type": "text"},
    {"question": "商业银行资本充足率不得低于多少？", "answer": "8%", "type": "text"},
    {"question": "商业银行杠杆率不得低于多少？", "answer": "4%", "type": "text"},
    {"question": "对一般企业债权的风险权重是多少？", "answer": "100%", "type": "text"},
    {"question": "银行业金融机构发生重大事项后，应当在多少日内报告？", "answer": "5个工作日", "type": "text"},
    {"question": "发生特别重大事项的，应当在多长时间内先行电话报告？", "answer": "24小时", "type": "text"},
    {"question": "个人消费贷款期限最长不得超过多少年？", "answer": "10年", "type": "text"},
    {"question": "单笔贷款金额超过多少万元的，应当追加有效担保？", "answer": "500万", "type": "text"},
    # ---- 多事实题 ----
    {"question": "根据资本管理办法，一级资本充足率和核心一级资本充足率分别不得低于多少？", "answer": "6%和5%", "type": "text"},
    # ---- 表格题：取数 ----
    {"question": "2024年一季度商业银行营业收入是多少亿元？", "answer": "135.2亿元", "type": "table"},
    {"question": "2024年三季度商业银行净利润是多少亿元？", "answer": "21.7亿元", "type": "table"},
    {"question": "2023年三季度商业银行不良贷款率是多少？", "answer": "1.79%", "type": "table"},
    {"question": "A机构2025年6月总资产是多少亿元？", "answer": "5200亿元", "type": "table"},
    {"question": "A机构2024年末营业收入是多少亿元？", "answer": "300亿元", "type": "table"},
    {"question": "B机构2024年末营业收入是多少亿元？", "answer": "450亿元", "type": "table"},
    # ---- 表格题：比较 / 计算（扩展） ----
    {"question": "B机构2024年末营业收入比A机构高多少亿元？", "answer": "150亿元", "type": "table"},
    {"question": "2024年一季度商业银行净利润是多少亿元？", "answer": "19.5亿元", "type": "table"},
    # ---- 无证据题 ----
    {"question": "某机构发生数据泄露后应当在多少日内报告金融监管总局？", "answer": "未找到足够证据", "type": "text"},
    {"question": "2020年一季度商业银行营业收入是多少亿元？", "answer": "未找到足够证据", "type": "table"},
]


def main() -> None:
    for d in (DOCS_DIR, TABLES_DIR, QUESTIONS_DIR, SAMPLE_DIR):
        d.mkdir(parents=True, exist_ok=True)

    # Word
    _write_docx(DOCS_DIR / "商业银行资本管理办法.docx", CAPITAL_RULES)
    _write_docx(DOCS_DIR / "银行业金融机构重大事项报告制度.docx", REPORT_RULES)

    # PDF
    _write_pdf(DOCS_DIR / "个人贷款管理办法.pdf", LOAN_RULES)

    # Excel
    _write_indicator_table(TABLES_DIR / "商业银行主要监管指标统计表.xlsx")
    _write_balance_table(TABLES_DIR / "银行业金融机构资产负债统计表.xlsx")

    # 测试集
    qpath = QUESTIONS_DIR / "testset.json"
    with open(qpath, "w", encoding="utf-8") as f:
        json.dump(QUESTIONS, f, ensure_ascii=False, indent=2)

    # 同步一份到 data/sample（保留生成源）
    for src in [DOCS_DIR / "商业银行资本管理办法.docx",
                DOCS_DIR / "银行业金融机构重大事项报告制度.docx",
                DOCS_DIR / "个人贷款管理办法.pdf"]:
        if src.exists():
            shutil.copy2(src, SAMPLE_DIR / src.name)
    for src in [TABLES_DIR / "商业银行主要监管指标统计表.xlsx",
                TABLES_DIR / "银行业金融机构资产负债统计表.xlsx"]:
        if src.exists():
            shutil.copy2(src, SAMPLE_DIR / src.name)
    if qpath.exists():
        shutil.copy2(qpath, SAMPLE_DIR / "testset.json")

    print(f"样例数据已生成：")
    print(f"  文档：{len(list(DOCS_DIR.glob('*'))) if DOCS_DIR.exists() else 0} 个")
    print(f"  表格：{len(list(TABLES_DIR.glob('*'))) if TABLES_DIR.exists() else 0} 个")
    print(f"  测试题：{len(QUESTIONS)} 条 -> {qpath}")


if __name__ == "__main__":
    main()
