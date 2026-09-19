"""Excel (.xlsx / .xlsm) 结构化解析器。

核心目标：把每个"数值单元格"与其上下文（行标签、列表头、单位、坐标）关联起来，
形成 :class:`~src.parsers.schemas.TableRecord` 列表，供后续表格取数。

设计要点：
- 使用 ``openpyxl.load_workbook(..., data_only=True)`` 读取公式的计算结果；
- 识别多级表头与合并单元格（forward fill）；
- 自动判定"表头行"与"标签列"，将每个数值单元格与行指标、列期间/机构关联；
- 检测单位（亿元 / 万元 / % 等）。

不做的事情：绝不把整个 Sheet 转成一段文本交给大模型猜测。
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Optional

import openpyxl
from openpyxl.utils import get_column_letter

from .schemas import TableRecord, cell_coordinate
from ..utils.config import TableConfig
from ..utils.logging import get_logger

log = get_logger("excel_parser")

# 常见单位关键字（按优先级）
_UNIT_PAT = re.compile(
    r"(\d+(?:\.\d+)?\s*(?:万亿|亿元|万元|千元|元|亿美元|美元|%|％|％|人|家|笔|户|次|天|日|年|月))"
)
_UNIT_TOKENS = ["万亿元", "亿元", "万元", "千元", "亿美元", "美元", "%", "％", "人", "家", "笔", "户", "次", "天", "日"]


def _is_number(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _is_label(v: Any) -> bool:
    return isinstance(v, str) and bool(v.strip())


def _parse_number(s: str) -> Optional[float]:
    """尝试把带千分位/百分号的字符串解析为 float。"""
    if not isinstance(s, str):
        return None
    t = s.strip().replace(",", "").replace("，", "")
    if t.endswith("%") or t.endswith("％"):
        t = t[:-1].strip()
        try:
            return float(t) / 100.0
        except ValueError:
            return None
    try:
        return float(t)
    except ValueError:
        return None


def _extract_unit(text: str) -> str:
    """从文本中提取单位。"""
    for tok in _UNIT_TOKENS:
        if tok in text:
            return tok
    return ""


def _strip_unit(text: str, unit: str) -> str:
    """去掉标签中的单位后缀，便于指标名匹配。"""
    if not unit:
        return text
    return text.replace(unit, "").strip(" （）()　")


def _load_grid(ws) -> dict[tuple[int, int], Any]:
    """读取 Sheet 的所有单元格值到 {(row, col): value}。

    对合并单元格做 forward fill：区域内空单元格继承左上角的值。
    """
    grid: dict[tuple[int, int], Any] = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                            min_col=1, max_col=ws.max_column):
        for cell in row:
            v = cell.value
            if v is not None:
                grid[(cell.row, cell.column)] = v

    # 合并单元格 forward fill
    for rng in ws.merged_cells.ranges:
        top_left = grid.get((rng.min_row, rng.min_col))
        if top_left is None:
            continue
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) not in grid:
                    grid[(r, c)] = top_left
    return grid


def _text_ratio(values: list[Any]) -> float:
    """非空值中字符串占比。"""
    non_empty = [v for v in values if v is not None and str(v).strip() != ""]
    if not non_empty:
        return 0.0
    labels = sum(1 for v in non_empty if _is_label(v))
    return labels / len(non_empty)


def _detect_header_rows(grid: dict, max_row: int, max_col: int) -> int:
    """返回表头区结束的行号（表头为 1..header_end，数据从 header_end+1 开始）。"""
    header_end = 0
    for r in range(1, max_row + 1):
        row_vals = [grid.get((r, c)) for c in range(1, max_col + 1)]
        if _text_ratio(row_vals) >= TableConfig.HEADER_TEXT_RATIO:
            header_end = r
        else:
            break
    # 限制表头最多 10 行，避免纯文本表被整体当作表头
    return min(header_end, 10)


def _detect_label_cols(grid: dict, data_start: int, max_row: int, max_col: int) -> int:
    """返回标签区结束的列号（标签为 1..label_end，数据列从 label_end+1 开始）。"""
    label_end = 0
    for c in range(1, max_col + 1):
        col_vals = [grid.get((r, c)) for r in range(data_start, max_row + 1)]
        if col_vals and _text_ratio(col_vals) >= TableConfig.HEADER_TEXT_RATIO:
            label_end = c
        else:
            break
    return label_end


def parse_excel(path: str | Path) -> tuple[list[TableRecord], list[str]]:
    """解析单个 Excel 文件，返回 (记录列表, 告警列表)。"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {path}")

    records: list[TableRecord] = []
    warnings: list[str] = []

    wb = openpyxl.load_workbook(str(path), data_only=True)
    for ws in wb.worksheets:
        if ws.max_row == 0 or ws.max_column == 0:
            warnings.append(f"Sheet「{ws.title}」为空")
            continue

        grid = _load_grid(ws)
        header_end = _detect_header_rows(grid, ws.max_row, ws.max_column)
        data_start = header_end + 1
        label_end = _detect_label_cols(grid, data_start, ws.max_row, ws.max_column)

        for r in range(data_start, ws.max_row + 1):
            for c in range(label_end + 1, ws.max_column + 1):
                raw = grid.get((r, c))
                if raw is None:
                    continue

                # 只保留"数值型"单元格：数字，或可解析为数字的字符串（如 "1,234.5"、"8%"）
                if isinstance(raw, str):
                    numeric = _parse_number(raw)
                    if numeric is None:
                        continue  # 数据区内的纯文本（备注等）跳过
                    value: Any = raw
                    meta_extra = {"numeric_value": numeric}
                elif _is_number(raw):
                    value = raw
                    meta_extra = {"numeric_value": float(raw)}
                else:
                    continue

                row_labels = [
                    str(grid[(r, j)]).strip()
                    for j in range(1, label_end + 1)
                    if (r, j) in grid and str(grid[(r, j)]).strip() != ""
                ]
                col_labels = [
                    str(grid[(i, c)]).strip()
                    for i in range(1, header_end + 1)
                    if (i, c) in grid and str(grid[(i, c)]).strip() != ""
                ]

                if not row_labels and not col_labels:
                    # 完全没有上下文标签的裸数值，跳过
                    continue

                # 单位：优先从列标签、其次行标签中检测
                unit = ""
                for lab in col_labels + row_labels:
                    unit = _extract_unit(lab)
                    if unit:
                        break
                # 清理标签中的单位
                row_labels = [_strip_unit(x, unit) for x in row_labels if _strip_unit(x, unit)]
                col_labels = [_strip_unit(x, unit) for x in col_labels if _strip_unit(x, unit)]

                records.append(
                    TableRecord(
                        source_file=path.name,
                        sheet=ws.title,
                        row=r,
                        column=c,
                        cell=cell_coordinate(r, c),
                        value=value,
                        row_header=" / ".join(row_labels),
                        column_header=" / ".join(col_labels),
                        row_labels=row_labels,
                        column_labels=col_labels,
                        headers={"unit": unit},
                        unit=unit,
                        metadata={"column_letter": get_column_letter(c), **meta_extra},
                    )
                )

    if not records:
        warnings.append(f"文件 {path.name} 未提取到任何带上下文的数值单元格")

    return records, warnings
